from django.shortcuts import render
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt, csrf_protect
from django.views.decorators.http import require_POST
import json
from openpyxl import Workbook, load_workbook
import os

KEY_PAIR_RESPONSES = {
    "hello": "Hello! How can I help you today?",
    "how are you": "I'm just a bot, but I'm doing great, thank you!",
    "what can you do": "I'm here to help you with any questions you have. Just ask me anything!",
    "who are you": "I'm a bot created by Mohammed Shaik Sahil. Nice to meet you!",
    "bye": "Goodbye! Have a great day!",
    "what is your name": "I'm a bot, you can call me Jack.",
    "help": "I'm here to help you! Just ask me anything."
}

SECRET_KEY = "admin123"

user_admin_state = {}
user_registration_state = {}

def save_to_excel(name, mobile):
    file_path = "register.xlsx"
    
    if not os.path.exists(file_path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Default Event"
        ws.append(['Name', 'Mobile'])
    else:
        wb = load_workbook(file_path)
        ws = wb.active
        event_name = ws.title

        # Check for duplicate registration
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == name and str(row[1]) == mobile:
                return "already_registered"

    ws.append([name, mobile])
    wb.save(file_path)
    return "registered"

@csrf_exempt
@require_POST
def get_bot_response(request):
    try:
        data = json.loads(request.body)
        user_message = data.get('message')
        user_id = data.get('user_id')

        if user_id not in user_admin_state:
            user_admin_state[user_id] = {
                'is_admin': False,
                'attempts': 0
            }
        if user_id not in user_registration_state:
            user_registration_state[user_id] = {
                'registration_step': 0,
                'name': '',
                'mobile': ''
            }

        if user_message.startswith("/admin"):
            if user_admin_state[user_id]['is_admin']:
                return JsonResponse({'response': "You are already in admin mode."})
            user_admin_state[user_id]['is_admin'] = True
            user_admin_state[user_id]['attempts'] = 0
            return JsonResponse({'response': "Please enter the secret key to enter admin mode."})
        
        elif user_message.startswith("/end") and user_admin_state[user_id]['is_admin']:
            user_admin_state[user_id]['is_admin'] = False
            return JsonResponse({'response': "You have exited admin mode."})
        
        if user_admin_state[user_id]['is_admin']:
            if user_message == SECRET_KEY and user_admin_state[user_id]['is_admin'] == True:
                user_admin_state[user_id]['is_admin'] = "granted"
                return JsonResponse({'response': "Hey admin, nice to meet you!"})
            return JsonResponse({'response': f"Admin mode: {user_message}"})
        
        if user_admin_state[user_id]['is_admin'] != "granted":
            registration_step = user_registration_state[user_id]['registration_step']

            if user_message.lower() == "register":
                user_registration_state[user_id]['registration_step'] = 1
                return JsonResponse({'response': "Please enter your name:"})

            if registration_step == 1:
                user_registration_state[user_id]['name'] = user_message
                user_registration_state[user_id]['registration_step'] = 2
                return JsonResponse({'response': "Please enter your mobile number:"})

            if registration_step == 2:
                if not user_message.isdigit() or len(user_message) != 10:
                    return JsonResponse({'response': "Please enter a valid 10-digit mobile number without any other characters."})

                user_registration_state[user_id]['mobile'] = user_message
                user_registration_state[user_id]['registration_step'] = 0
                user_name = user_registration_state[user_id]['name']
                user_mobile = user_registration_state[user_id]['mobile']
                registration_status = save_to_excel(user_name, user_mobile)

                if registration_status == "already_registered":
                    return JsonResponse({'response': "You are already registered with this name and mobile number."})
                else:
                    del user_registration_state[user_id]
                    try:
                        wb = load_workbook("register.xlsx")
                        ws = wb.active
                        event_name = ws.title
                    except:
                        event_name = "Default Event"
                    return JsonResponse({'response': f"User registered for event {event_name} with name: {user_name} and mobile: {user_mobile}."})
        else:
            return JsonResponse({'response': "You are an admin and cannot perform user registration."})

        user_message = user_message.strip().lower()
        bot_response = KEY_PAIR_RESPONSES.get(user_message, "Sorry, I couldn't understand that. Could you please rephrase?")
        return JsonResponse({'response': bot_response})

    except Exception as e:
        return JsonResponse({'error': 'An error occurred while processing your request. Please try again later.'}, status=500)

def index(request):
    return render(request, 'index.html')

def message(request):
    return render(request, 'message.html')
